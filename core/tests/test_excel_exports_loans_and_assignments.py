"""
Tests for two new Excel exports:
1. /loans/?export=excel — full loan register, respecting active filters,
   with every vital field (not just what fits in the on-screen table).
2. /assignments/<id>/?export=excel (and the review page) — every client
   affected by an assignment request, human-readable, not just raw IDs.
"""
from decimal import Decimal
from datetime import timedelta

import openpyxl
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from io import BytesIO

from core.models import AssignmentRequest, Loan, LoanRepaymentSchedule
from core.tests.factories import make_branch, make_user, make_client, make_loan_product

EXCEL_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def read_workbook(response):
    return openpyxl.load_workbook(BytesIO(response.content))


class LoanListExcelExportTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.branch = make_branch(code='XLL001')
        cls.other_branch = make_branch(name='XLL001 Other', code='XLL001B')
        cls.director = make_user(cls.branch, role='director', email='xll_dir@test.com')
        cls.staff = make_user(cls.branch, role='staff', email='xll_staff@test.com')
        cls.client_obj = make_client(cls.branch, cls.staff, email='xll_client@test.com')
        cls.product = make_loan_product(code='XLL001P', loan_type='regular')

        cls.active_loan = Loan.objects.create(
            client=cls.client_obj, loan_product=cls.product, branch=cls.branch,
            principal_amount=Decimal('300000.00'), duration_months=6,
            disbursement_method='cash', created_by=cls.staff,
            purpose='Trading', status='active',
            outstanding_balance=Decimal('200000.00'), amount_paid=Decimal('100000.00'),
            disbursement_date=timezone.now(), application_date=timezone.now(),
        )
        cls.completed_loan = Loan.objects.create(
            client=cls.client_obj, loan_product=cls.product, branch=cls.branch,
            principal_amount=Decimal('50000.00'), duration_months=3,
            disbursement_method='cash', created_by=cls.staff,
            purpose='Farming', status='completed',
            outstanding_balance=Decimal('0.00'), amount_paid=Decimal('55000.00'),
            disbursement_date=timezone.now(), application_date=timezone.now(),
        )

    def setUp(self):
        self.client.force_login(self.director)

    def test_export_returns_excel_content_type(self):
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], EXCEL_CONTENT_TYPE)

    def test_export_contains_all_vital_columns(self):
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel'})
        wb = read_workbook(response)
        ws = wb['Loans']
        header_row = [c.value for c in ws[3]]  # row 1=title, 2=subtitle, 3=headers
        for expected in (
            'Loan Number', 'Client ID', 'Client Name', 'Branch', 'Loan Officer',
            'Product', 'Status', 'Principal (₦)', 'Outstanding Balance (₦)',
            'Amount Paid (₦)', 'Total Interest (₦)', 'Total Repayment (₦)',
            'Application Date', 'Disbursement Date', 'Purpose',
        ):
            self.assertIn(expected, header_row, f'missing column: {expected}')

    def test_export_includes_every_loan_matching_no_filter(self):
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel'})
        wb = read_workbook(response)
        ws = wb['Loans']
        loan_numbers = {row[0].value for row in ws.iter_rows(min_row=4) if row[0].value}
        self.assertEqual(loan_numbers, {self.active_loan.loan_number, self.completed_loan.loan_number})

    def test_export_respects_status_filter(self):
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel', 'status': 'active'})
        wb = read_workbook(response)
        ws = wb['Loans']
        loan_numbers = {row[0].value for row in ws.iter_rows(min_row=4) if row[0].value}
        self.assertEqual(loan_numbers, {self.active_loan.loan_number})

    def test_staff_export_is_scoped_to_own_clients(self):
        other_client = make_client(self.other_branch, make_user(self.other_branch, role='staff', email='xll_other_staff@test.com'), email='xll_other_client@test.com')
        Loan.objects.create(
            client=other_client, loan_product=self.product, branch=self.other_branch,
            principal_amount=Decimal('10000.00'), duration_months=1,
            disbursement_method='cash', created_by=other_client.assigned_staff,
            purpose='Other', status='active', disbursement_date=timezone.now(), application_date=timezone.now(),
        )
        self.client.force_login(self.staff)
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel'})
        wb = read_workbook(response)
        ws = wb['Loans']
        loan_numbers = {row[0].value for row in ws.iter_rows(min_row=4) if row[0].value}
        self.assertEqual(loan_numbers, {self.active_loan.loan_number, self.completed_loan.loan_number})


class LoanListExcelExportArrearsColumnsTests(TestCase):
    """The specific regulatory-style aging columns requested: Customer name,
    Date of disbursement, Loan product, Amount disbursed, Repayment amount,
    Amount paid already, Overdue (P), Total overdue, Balance Default, the
    five day-range buckets, and Days in Arrears."""

    @classmethod
    def setUpTestData(cls):
        cls.branch = make_branch(code='XLL002')
        cls.director = make_user(cls.branch, role='director', email='xll2_dir@test.com')
        cls.staff = make_user(cls.branch, role='staff', email='xll2_staff@test.com')
        cls.client_obj = make_client(cls.branch, cls.staff, email='xll2_client@test.com')
        cls.product = make_loan_product(code='XLL002P', loan_type='regular')

        cls.loan = Loan.objects.create(
            client=cls.client_obj, loan_product=cls.product, branch=cls.branch,
            principal_amount=Decimal('120000.00'), duration_months=6,
            disbursement_method='cash', created_by=cls.staff,
            purpose='Trading', status='active',
            outstanding_balance=Decimal('80000.00'), amount_paid=Decimal('40000.00'),
            amount_disbursed=Decimal('120000.00'),
            disbursement_date=timezone.now(), application_date=timezone.now(),
        )
        today = timezone.now().date()
        # One paid installment (must NOT count toward arrears).
        LoanRepaymentSchedule.objects.create(
            loan=cls.loan, installment_number=1, due_date=today - timedelta(days=90),
            principal_amount=Decimal('20000.00'), interest_amount=Decimal('5000.00'),
            total_amount=Decimal('25000.00'), amount_paid=Decimal('25000.00'),
            outstanding_amount=Decimal('0.00'), status='paid',
        )
        # Oldest unpaid overdue installment -> 45 days overdue -> bucket 31-60.
        LoanRepaymentSchedule.objects.create(
            loan=cls.loan, installment_number=2, due_date=today - timedelta(days=45),
            principal_amount=Decimal('20000.00'), interest_amount=Decimal('5000.00'),
            total_amount=Decimal('25000.00'), amount_paid=Decimal('0.00'),
            outstanding_amount=Decimal('25000.00'), status='pending',
        )
        # A second unpaid overdue installment, more recent -> still counts
        # toward the totals, but doesn't change "oldest" / days-in-arrears.
        LoanRepaymentSchedule.objects.create(
            loan=cls.loan, installment_number=3, due_date=today - timedelta(days=10),
            principal_amount=Decimal('20000.00'), interest_amount=Decimal('5000.00'),
            total_amount=Decimal('25000.00'), amount_paid=Decimal('0.00'),
            outstanding_amount=Decimal('25000.00'), status='pending',
        )
        # Not yet due -> must not count as overdue at all.
        LoanRepaymentSchedule.objects.create(
            loan=cls.loan, installment_number=4, due_date=today + timedelta(days=20),
            principal_amount=Decimal('20000.00'), interest_amount=Decimal('5000.00'),
            total_amount=Decimal('25000.00'), amount_paid=Decimal('0.00'),
            outstanding_amount=Decimal('25000.00'), status='pending',
        )

    def setUp(self):
        self.client.force_login(self.director)
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel'})
        wb = read_workbook(response)
        self.ws = wb['Loans']
        header = [c.value for c in self.ws[3]]
        self.col = {name: i for i, name in enumerate(header)}
        self.row = [c.value for c in self.ws[4]]

    def _get(self, col_name):
        return self.row[self.col[col_name]]

    def test_all_requested_columns_are_present(self):
        for expected in (
            'Customer name', 'Date of disbursement', 'Loan product', 'Amount disbursed (₦)',
            'Repayment amount (₦)', 'Amount paid already (₦)', 'Overdue (P) (₦)',
            'Total overdue (₦)', 'Balance Default (what has not been paid) (₦)',
            '1 to 30 (₦)', '31 to 60 (₦)', '61 to 90 (₦)', '91 to 180 (₦)', '181 to 360 (₦)',
            'Days in Arrears',
        ):
            self.assertIn(expected, self.col, f'missing column: {expected}')

    def test_customer_name_and_loan_product_and_disbursement_date(self):
        self.assertEqual(self._get('Customer name'), self.client_obj.get_full_name())
        self.assertEqual(self._get('Loan product'), self.product.name)
        self.assertEqual(self._get('Date of disbursement'), self.loan.disbursement_date.strftime('%Y-%m-%d'))

    def test_amount_disbursed_and_repayment_and_paid(self):
        self.assertEqual(self._get('Amount disbursed (₦)'), 120000.0)
        self.assertEqual(self._get('Amount paid already (₦)'), 40000.0)

    def test_overdue_principal_only_counts_unpaid_overdue_installments(self):
        # Two unpaid overdue installments of 20,000 principal each = 40,000.
        # The paid one and the not-yet-due one must be excluded.
        self.assertEqual(self._get('Overdue (P) (₦)'), 40000.0)

    def test_total_overdue_sums_outstanding_of_unpaid_overdue_installments(self):
        # Two unpaid overdue installments of 25,000 outstanding each = 50,000.
        self.assertEqual(self._get('Total overdue (₦)'), 50000.0)

    def test_balance_default_is_the_whole_loan_outstanding_balance(self):
        self.assertEqual(self._get('Balance Default (what has not been paid) (₦)'), 80000.0)

    def test_days_in_arrears_uses_the_oldest_unpaid_overdue_installment(self):
        self.assertEqual(self._get('Days in Arrears'), 45)

    def test_total_overdue_falls_in_the_31_to_60_bucket_only(self):
        self.assertEqual(self._get('1 to 30 (₦)'), 0.0)
        self.assertEqual(self._get('31 to 60 (₦)'), 50000.0)
        self.assertEqual(self._get('61 to 90 (₦)'), 0.0)
        self.assertEqual(self._get('91 to 180 (₦)'), 0.0)
        self.assertEqual(self._get('181 to 360 (₦)'), 0.0)

    def test_current_loan_with_no_overdue_installments_has_zero_arrears(self):
        current_loan = Loan.objects.create(
            client=self.client_obj, loan_product=self.product, branch=self.branch,
            principal_amount=Decimal('50000.00'), duration_months=6,
            disbursement_method='cash', created_by=self.staff,
            purpose='Trading', status='active',
            outstanding_balance=Decimal('50000.00'), amount_disbursed=Decimal('50000.00'),
            disbursement_date=timezone.now(), application_date=timezone.now(),
        )
        response = self.client.get(reverse('core:loan_list'), {'export': 'excel'})
        wb = read_workbook(response)
        ws = wb['Loans']
        header = [c.value for c in ws[3]]
        col = {name: i for i, name in enumerate(header)}
        for row in ws.iter_rows(min_row=4):
            if row[col['Loan Number']].value == current_loan.loan_number:
                self.assertEqual(row[col['Days in Arrears']].value, 0)
                for bucket in ('1 to 30 (₦)', '31 to 60 (₦)', '61 to 90 (₦)', '91 to 180 (₦)', '181 to 360 (₦)'):
                    self.assertEqual(row[col[bucket]].value, 0.0)
                break
        else:
            self.fail('current_loan not found in export')


class AssignmentAffectedClientsExcelExportTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.branch = make_branch(code='XLA001')
        cls.manager = make_user(cls.branch, role='manager', email='xla_mgr@test.com')
        cls.staff_a = make_user(cls.branch, role='staff', email='xla_a@test.com')
        cls.staff_b = make_user(cls.branch, role='staff', email='xla_b@test.com')
        cls.clients = [make_client(cls.branch, cls.staff_a, email=f'xla_c{i}@test.com') for i in range(3)]
        cls.req = AssignmentRequest.objects.create(
            assignment_type='bulk_clients_to_staff',
            assignment_data={'client_ids': [str(c.id) for c in cls.clients], 'staff_id': str(cls.staff_b.id)},
            description='Transfer 3 clients from Staff A to Staff B',
            requested_by=cls.staff_a,
            branch=cls.branch,
            affected_count=3,
        )

    def setUp(self):
        self.client.force_login(self.manager)

    def test_detail_page_export_returns_excel(self):
        response = self.client.get(reverse('core:assignment_detail', args=[self.req.id]), {'export': 'excel'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], EXCEL_CONTENT_TYPE)

    def test_detail_page_export_contains_every_client(self):
        response = self.client.get(reverse('core:assignment_detail', args=[self.req.id]), {'export': 'excel'})
        wb = read_workbook(response)
        ws = wb['Affected Clients']
        names = {row[1].value for row in ws.iter_rows(min_row=4) if row[1].value}
        self.assertEqual(names, {c.get_full_name() for c in self.clients})

    def test_review_page_export_also_works(self):
        response = self.client.get(reverse('core:assignment_approve', args=[self.req.id]), {'export': 'excel'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], EXCEL_CONTENT_TYPE)

    def test_export_button_present_on_detail_page(self):
        response = self.client.get(reverse('core:assignment_detail', args=[self.req.id]))
        self.assertContains(response, 'Export to Excel')

    def test_unauthorized_user_cannot_export(self):
        outsider = make_user(self.branch, role='staff', email='xla_outsider@test.com')
        self.client.force_login(outsider)
        response = self.client.get(reverse('core:assignment_detail', args=[self.req.id]), {'export': 'excel'})
        self.assertEqual(response.status_code, 403)
